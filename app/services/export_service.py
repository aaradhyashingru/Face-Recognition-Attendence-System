import io
from datetime import datetime
from typing import List
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from app.models import AttendanceLog


class ExportService:
    @staticmethod
    def generate_excel(logs: List[AttendanceLog]) -> io.BytesIO:
        """Generate a professionally styled Excel workbook of attendance records."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance Log"

        # Styles
        title_font = Font(name="Calibri", size=16, bold=True, color="1F2937")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="E5E7EB"),
            right=Side(style="thin", color="E5E7EB"),
            top=Side(style="thin", color="E5E7EB"),
            bottom=Side(style="thin", color="E5E7EB"),
        )

        # Title Block
        ws.merge_cells("A1:H1")
        ws["A1"] = "FACE RECOGNITION ATTENDANCE SYSTEM — AUDIT REPORT"
        ws["A1"].font = title_font
        ws["A1"].alignment = center_align

        ws.merge_cells("A2:H2")
        ws["A2"] = f"Generated On: {datetime.now().strftime('%d %B %Y, %I:%M %p')} | Total Logs: {len(logs)}"
        ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="6B7280")
        ws["A2"].alignment = center_align

        # Table Headers
        headers = [
            "#",
            "User Code",
            "Full Name",
            "Department",
            "Date",
            "Time",
            "Punch Type",
            "Status",
        ]
        ws.append([])  # blank row 3
        ws.append(headers)  # row 4

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        # Table Rows
        for i, log in enumerate(logs, 1):
            dept_name = log.user.department.name if log.user and log.user.department else "General"
            row_data = [
                i,
                log.user.user_code if log.user else "N/A",
                log.user.full_name if log.user else "N/A",
                dept_name,
                log.timestamp.strftime("%Y-%m-%d"),
                log.timestamp.strftime("%I:%M:%S %p"),
                log.punch_type,
                log.status,
            ]
            ws.append(row_data)
            row_idx = ws.max_row

            # Apply row formatting
            for col_num in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.border = thin_border
                cell.alignment = center_align if col_num in [1, 2, 5, 6, 7, 8] else left_align
                if i % 2 == 0:
                    cell.fill = alt_fill

        # Auto-fit Column Widths
        for col_idx, col in enumerate(ws.columns, 1):
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def generate_pdf(logs: List[AttendanceLog]) -> io.BytesIO:
        """Generate a landscape PDF report of attendance records."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(letter),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            name="TitleStyle",
            parent=styles["Heading1"],
            fontSize=18,
            textColor=colors.HexColor("#1F2937"),
            alignment=1,  # Center
            spaceAfter=5,
        )
        subtitle_style = ParagraphStyle(
            name="SubtitleStyle",
            parent=styles["Normal"],
            fontSize=10,
            textColor=colors.HexColor("#6B7280"),
            alignment=1,
            spaceAfter=15,
        )

        elements = [
            Paragraph("Face Recognition Attendance System — Audit Report", title_style),
            Paragraph(
                f"Generated on {datetime.now().strftime('%d %b %Y, %I:%M %p')} | Total Logs: {len(logs)}",
                subtitle_style,
            ),
            Spacer(1, 10),
        ]

        table_data = [
            ["#", "User Code", "Full Name", "Department", "Date", "Time", "Type", "Status"]
        ]

        for i, log in enumerate(logs, 1):
            dept_name = log.user.department.name if log.user and log.user.department else "General"
            table_data.append([
                str(i),
                log.user.user_code if log.user else "N/A",
                log.user.full_name if log.user else "N/A",
                dept_name,
                log.timestamp.strftime("%Y-%m-%d"),
                log.timestamp.strftime("%I:%M:%S %p"),
                log.punch_type,
                log.status,
            ])

        pdf_table = Table(table_data, repeatRows=1)
        pdf_table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3B82F6")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 10),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 6),
                ("TOPPADDING", (0, 0), (-1, 0), 6),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 1), (-1, -1), 9),
            ])
        )

        elements.append(pdf_table)
        doc.build(elements)
        output.seek(0)
        return output
